import hashlib
import html
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
import streamlit as st
from azure.ai.documentintelligence import DocumentIntelligenceClient
from azure.ai.documentintelligence.models import (
    AnalyzeDocumentRequest,
    DocumentAnalysisFeature,
)
from azure.core.credentials import AzureKeyCredential
from azure.core.exceptions import HttpResponseError
from dotenv import load_dotenv

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

# Optional, used only to preview Office files:  pip install mammoth python-pptx openpyxl
try:
    import mammoth
except ImportError:
    mammoth = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

load_dotenv(Path(__file__).resolve().parent / ".env", override=True)


def env(*names: str) -> str:
    for n in names:
        if v := os.getenv(n, "").strip().strip('"').strip("'"):
            return v
    return ""


ENV_ENDPOINT = env("DOCUMENTINTELLIGENCE_ENDPOINT", "AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT")
ENV_KEY = env("DOCUMENTINTELLIGENCE_API_KEY", "AZURE_DOCUMENT_INTELLIGENCE_KEY")

st.set_page_config(page_title="Document Intelligence", page_icon="📄", layout="wide")

# ----------------------------------------------------------------------------
# The 5 prebuilt models. `tabs` = which result tables are shown for that model.
# ----------------------------------------------------------------------------
MODELS = {
    "Prebuilt Invoice": {
        "model_id": "prebuilt-invoice",
        "features": [],
        "office": False,  # PDF / images only
        "tabs": ["Fields", "Line items"],
        "help": "Vendor, customer, dates, totals, tax and line items.",
    },
    "Prebuilt Receipt": {
        "model_id": "prebuilt-receipt",
        "features": [],
        "office": False,  # PDF / images only
        "tabs": ["Fields", "Line items"],
        "help": "Merchant, date, totals, tax, tip and purchased items.",
    },
    "Prebuilt Read": {
        "model_id": "prebuilt-read",
        "features": [],
        "office": True,
        "tabs": ["Lines", "Paragraphs", "Words", "Languages"],
        "help": "OCR: printed and handwritten text, words with confidence, detected languages.",
    },
    "Prebuilt Layout": {
        "model_id": "prebuilt-layout",
        "features": [],  # key-value pairs are not shown in this view, and add-ons do not work on Office files
        "office": True,
        "tabs": ["Paragraphs", "Tables", "Selection marks"],
        "help": "Document structure: paragraphs with roles, tables and selection marks.",
    },
    "Prebuilt General Document": {
        "model_id": "prebuilt-layout",
        "features": [DocumentAnalysisFeature.KEY_VALUE_PAIRS],
        "office": False,  # the key-value add-on is not supported on Office / HTML files
        "tabs": ["Key-value pairs", "Tables", "Paragraphs"],
        "help": (
            "Key-value pairs, tables and text. In API v4 the old prebuilt-document model "
            "is replaced by Layout + the key-value-pairs add-on, which is what runs here. PDF and images only."
        ),
    },
}
ALLOWED_TYPES = ["pdf", "jpg", "jpeg", "png", "bmp", "tiff", "tif", "heif", "docx", "xlsx", "pptx", "html"]
OFFICE_TYPES = {"docx", "xlsx", "pptx", "html"}  # only Read and Layout accept these, and without add-on features
NO_PREVIEW_TYPES = OFFICE_TYPES
MAX_WORKERS = 4
STEPS = ["Choose model", "Add documents", "Analyze", "Review results"]


# ----------------------------------------------------------------------------
# Azure helpers
# ----------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def get_client(endpoint: str, key: str) -> DocumentIntelligenceClient:
    return DocumentIntelligenceClient(endpoint=endpoint, credential=AzureKeyCredential(key))


def analyze(client: DocumentIntelligenceClient, cfg: dict, data: bytes, filename: str):
    is_office = ext_of(filename) in OFFICE_TYPES
    if is_office and not cfg["office"]:
        raise ValueError(f"{cfg['model_id']} accepts PDF and images only, not .{ext_of(filename)} files.")
    features = [] if is_office else cfg["features"]  # add-on capabilities are not supported for Office files
    kwargs = {"features": features} if features else {}
    poller = client.begin_analyze_document(
        cfg["model_id"], body=AnalyzeDocumentRequest(bytes_source=data), **kwargs
    )
    return poller.result()


def _conf(x):
    return round(x, 3) if x is not None else None


def _val(x):
    # converts an enum to its plain value
    return getattr(x, "value", x)


def table_to_df(t) -> pd.DataFrame:
    grid = [[""] * t.column_count for _ in range(t.row_count)]
    for c in t.cells:
        grid[c.row_index][c.column_index] = c.content or ""
    has_header = any(c.row_index == 0 and c.kind == "columnHeader" for c in t.cells)
    if has_header and len(grid) > 1:
        cols, seen = [], {}
        for i, h in enumerate(grid[0]):
            h = h.strip() or f"Column {i + 1}"
            seen[h] = seen.get(h, 0) + 1
            cols.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
        return pd.DataFrame(grid[1:], columns=cols)
    return pd.DataFrame(grid)


def parse(result) -> dict:
    """Turn the SDK result into plain Python / pandas objects (safe to keep in session state)."""
    docs = result.documents or []
    multi = len(docs) > 1

    # --- prebuilt models with fields (invoice, receipt) ---
    fields, arrays, doc_types = [], {}, []
    for i, doc in enumerate(docs, 1):
        if doc.doc_type:
            doc_types.append(f"{doc.doc_type} ({_conf(doc.confidence)})")
        base = {"Doc": i} if multi else {}
        for name, f in (doc.fields or {}).items():
            if f.value_array:  # e.g. Items on invoices / receipts
                for item in f.value_array:
                    row = dict(base)
                    row.update({k: (v.content if v else None) for k, v in (item.value_object or {}).items()})
                    arrays.setdefault(name, []).append(row)
            else:
                fields.append({**base, "Field": name, "Value": f.content, "Confidence": _conf(f.confidence)})

    # --- pages: lines, words, selection marks ---
    lines, words, marks = [], [], []
    for page in result.pages or []:
        pn = page.page_number
        for i, ln in enumerate(page.lines or [], 1):
            lines.append({"Page": pn, "Line": i, "Text": ln.content})
        for w in page.words or []:
            words.append({"Page": pn, "Word": w.content, "Confidence": _conf(w.confidence)})
        for m in page.selection_marks or []:
            marks.append({"Page": pn, "State": _val(m.state), "Confidence": _conf(m.confidence)})

    # --- paragraphs, languages, key-value pairs, tables ---
    paragraphs = []
    for p in result.paragraphs or []:
        page_no = p.bounding_regions[0].page_number if p.bounding_regions else None
        paragraphs.append({"Page": page_no, "Role": _val(p.role) or "text", "Text": p.content})

    lang_rows = [
        {
            "Language": lg.locale,
            "Confidence": _conf(lg.confidence),
            "Characters": sum(s.length for s in (lg.spans or [])),
        }
        for lg in result.languages or []
    ]
    languages = pd.DataFrame(lang_rows)
    if not languages.empty:
        languages = languages.groupby("Language", as_index=False).agg( Confidence=("Confidence", "max"), Characters=("Characters", "sum") )

    kv = [
        {
            "Key": p.key.content if p.key else "",
            "Value": p.value.content if p.value else "",
            "Confidence": _conf(p.confidence),
        }
        for p in result.key_value_pairs or []
    ]

    tables = []
    for t in result.tables or []:
        page_no = t.bounding_regions[0].page_number if t.bounding_regions else "?"
        tables.append((f"Table {len(tables) + 1} · page {page_no}", table_to_df(t)))

    def one(rows):
        df = rows if isinstance(rows, pd.DataFrame) else pd.DataFrame(rows)
        return [("", df)] if not df.empty else []

    sections = {
        "Fields": one(fields),
        "Line items": [(k, pd.DataFrame(v)) for k, v in arrays.items()],
        "Key-value pairs": one(kv),
        "Tables": tables,
        "Paragraphs": one(paragraphs),
        "Lines": one(lines),
        "Words": one(words),
        "Languages": one(languages),
        "Selection marks": one(marks),
    }

    stats = [f"{len(result.pages or [])} page(s)"]
    if doc_types:
        stats.append("detected: " + ", ".join(doc_types))
    if any(s.is_handwritten for s in (result.styles or [])):
        stats.append("handwriting detected")

    try:
        raw = result.as_dict()
    except AttributeError:
        raw = dict(result)

    return {
        "sections": sections,
        "stats": stats,
        "text": result.content or "",
        "json": json.dumps(raw, indent=2, default=str),
    }


# ----------------------------------------------------------------------------
# Styling
# ----------------------------------------------------------------------------
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Public+Sans:wght@400;500;600;700&display=swap');

:root {
  --bg: #EEF1F6;
  --surface: #FFFFFF;
  --ink: #172033;
  --muted: #5B6577;
  --line: #D5DBE5;
  --accent: #2F3FD0;
  --accent-soft: #F4F5FF;
  --ok: #1E7A4F;
  --err: #B3261E;
}

.stApp,
.stApp :where(p, span, div, label, button, input, textarea, li, a, h1, h2, h3, h4, code):not([data-testid="stIconMaterial"]) {
  font-family: 'Public Sans', system-ui, -apple-system, 'Segoe UI', sans-serif;
}
.stApp { background: var(--bg); color: var(--ink); }
.block-container { padding-top: 2rem; padding-bottom: 7rem; max-width: 1500px; }
header[data-testid="stHeader"] { background: transparent; }

/* title */
.app-title { font-size: 1.6rem; font-weight: 700; letter-spacing: -0.01em; margin: 0; }
.app-sub { color: var(--muted); margin: .15rem 0 1.25rem; font-size: .95rem; }

/* step tracker */
.steps { display: flex; align-items: center; margin: 0 0 1.5rem; }
.step { display: flex; align-items: center; gap: .6rem; color: var(--muted); font-weight: 500; font-size: .92rem; white-space: nowrap; }
.step:not(:last-child) { flex: 1; }
.step:not(:last-child)::after { content: ""; flex: 1; height: 2px; background: var(--line); margin: 0 .9rem; }
.step .dot { width: 1.7rem; height: 1.7rem; border-radius: 50%; display: grid; place-items: center;
             border: 2px solid var(--line); background: var(--surface); font-size: .8rem; font-weight: 700; flex: none; }
.step.done { color: var(--ink); }
.step.done .dot { background: var(--ok); border-color: var(--ok); color: #fff; }
.step.done:not(:last-child)::after { background: var(--ok); }
.step.active { color: var(--ink); font-weight: 700; }
.step.active .dot { background: var(--accent); border-color: var(--accent); color: #fff; box-shadow: 0 0 0 4px rgba(47,63,208,.15); }
@media (max-width: 800px) { .step .lbl { display: none; } }

/* panels */
.st-key-controls, .st-key-workspace { background: var(--surface); border: 1px solid var(--line); border-radius: 14px; padding: 1.25rem; }
.panel-h { font-weight: 700; font-size: 1rem; margin: 0 0 .5rem; }
.hint { color: var(--muted); font-size: .85rem; margin: .35rem 0 0; }

/* drop zone */
[data-testid="stFileUploaderDropzone"] {
  border: 2px dashed var(--accent); background: var(--accent-soft); border-radius: 14px;
  padding: 1.75rem 1rem; flex-direction: column; text-align: center; gap: .75rem;
}
[data-testid="stFileUploaderDropzoneInstructions"] { flex-direction: column; align-items: center; }
[data-testid="stFileUploaderDropzoneInstructions"] small { display: none; }
[data-testid="stFileUploaderDropzone"]:hover { background: #ECEEFF; }

/* format pills */
.pills { display: flex; flex-wrap: wrap; gap: .35rem; margin: .6rem 0 .25rem; }
.pill { background: #E7EAF2; color: var(--ink); font-size: .72rem; font-weight: 600; padding: .15rem .55rem; border-radius: 999px; }
.pill.limit { background: var(--ink); color: #fff; }

/* file queue */
.qname { font-weight: 600; font-size: .9rem; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
.qmeta { font-size: .78rem; color: var(--muted); }
.bar { height: 5px; background: var(--line); border-radius: 3px; overflow: hidden; margin-top: .35rem; }
.bar i { display: block; height: 100%; width: 100%; background: var(--accent); }
.bar i.done { background: var(--ok); }
.bar i.err { background: var(--err); }
.ext { width: 44px; height: 56px; border-radius: 6px; background: #E7EAF2; display: grid; place-items: center;
       font-size: .68rem; font-weight: 700; color: var(--muted); }

/* results */
.chips { display: flex; flex-wrap: wrap; gap: .4rem; margin: 0 0 .75rem; }
.chip { background: var(--accent-soft); color: var(--accent); border: 1px solid #D9DCFA; font-size: .78rem; font-weight: 600; padding: .15rem .6rem; border-radius: 999px; }
.empty { border: 1px dashed var(--line); border-radius: 12px; padding: 2.5rem 1.5rem; text-align: center; color: var(--muted); }
.empty b { display: block; color: var(--ink); font-size: 1.05rem; margin-bottom: .35rem; }

/* floating call-to-action */
.st-key-cta { position: fixed; right: 2rem; bottom: 1.5rem; z-index: 1000; width: auto !important; }
.st-key-cta button { border-radius: 999px; padding: .75rem 1.6rem; font-weight: 700;
                     box-shadow: 0 8px 24px rgba(23,32,51,.25); }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def fmt_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024


def ext_of(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower()


def stepper_html(current: int) -> str:
    """current is 1-based; steps before it are shown as done."""
    parts = []
    for i, label in enumerate(STEPS, 1):
        state = "done" if i < current else "active" if i == current else ""
        dot = "✓" if i < current else str(i)
        parts.append(f'<div class="step {state}"><span class="dot">{dot}</span><span class="lbl">{label}</span></div>')
    return f'<div class="steps">{"".join(parts)}</div>'


@st.cache_data(show_spinner=False, max_entries=32)
def page_images(doc_id: str, filename: str, _data: bytes, max_pages: int = 10):
    """Render the first pages of a PDF / image to PNG bytes. `doc_id` is the cache key."""
    doc = fitz.open(stream=_data, filetype=ext_of(filename))
    try:
        total = doc.page_count
        imgs = [doc[i].get_pixmap(dpi=110).tobytes("png") for i in range(min(total, max_pages))]
    finally:
        doc.close()
    return imgs, total


@st.cache_data(show_spinner=False, max_entries=64)
def thumbnail(doc_id: str, filename: str, _data: bytes):
    """First page as a small PNG, or None when it can't be rendered."""
    if fitz is None or ext_of(filename) in NO_PREVIEW_TYPES:
        return None
    try:
        doc = fitz.open(stream=_data, filetype=ext_of(filename))
        try:
            return doc[0].get_pixmap(dpi=36).tobytes("png")
        finally:
            doc.close()
    except Exception:  # noqa: BLE001
        return None


def _soffice_path():
    for name in ("soffice", "libreoffice"):
        if p := shutil.which(name):
            return p
    for p in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(p):
            return p
    return None


@st.cache_data(show_spinner="Rendering preview...", max_entries=16)
def office_to_pdf(doc_id: str, filename: str, _data: bytes):
    """Word / Excel / PowerPoint -> PDF bytes via LibreOffice, or None when it isn't installed or fails."""
    exe = _soffice_path()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"doc.{ext_of(filename)}"
        src.write_bytes(_data)
        profile = (Path(tmp) / "profile").as_uri()  # private profile: no clash with an open LibreOffice
        try:
            subprocess.run(
                [exe, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                 "--outdir", tmp, str(src)],
                check=True, capture_output=True, timeout=90,
            )
        except Exception:  # noqa: BLE001
            return None
        pdf = Path(tmp) / "doc.pdf"
        return pdf.read_bytes() if pdf.exists() else None


def office_html(filename: str, data: bytes):
    """Simplified HTML rendering of an Office / HTML file with pure-Python libraries, or None."""
    ext = ext_of(filename)
    try:
        if ext == "html":
            text = data.decode("utf-8", errors="replace")
            if m := re.search(r"(?is)<body[^>]*>(.*)</body>", text):
                text = m.group(1)
            # scripts / styles from an uploaded file must not touch the app
            return re.sub(r"(?is)<(script|style)\b.*?</\1>|<link\b[^>]*>", "", text)
        if ext == "docx" and mammoth:
            return mammoth.convert_to_html(io.BytesIO(data)).value
        if ext == "xlsx":
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
            return "".join(
                f"<h3>{html.escape(str(name))}</h3>" + df.fillna("").to_html(header=False, index=False, border=1)
                for name, df in sheets.items()
            )
        if ext == "pptx" and Presentation:
            out = []
            for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
                out.append(f"<h3>Slide {i}</h3>")
                for sh in slide.shapes:
                    if getattr(sh, "has_table", False) and sh.has_table:
                        rows = "".join(
                            "<tr>" + "".join(f"<td>{html.escape(c.text)}</td>" for c in r.cells) + "</tr>"
                            for r in sh.table.rows
                        )
                        out.append(f"<table>{rows}</table>")
                    elif sh.has_text_frame and sh.text_frame.text.strip():
                        out.append("<p>" + html.escape(sh.text_frame.text).replace("\n", "<br>") + "</p>")
            return "".join(out)
    except Exception:  # noqa: BLE001
        return None
    return None


_PREVIEW_CSS = (
    ".office-preview{font-size:14px;color:#172033;}"
    ".office-preview table{border-collapse:collapse;margin:8px 0;}"
    ".office-preview td,.office-preview th{border:1px solid #D5DBE5;padding:4px 8px;vertical-align:top;}"
    ".office-preview :is(p,td,th,li,h1,h2,h3,h4,h5,h6){unicode-bidi:plaintext;}"  # each paragraph picks its own direction (Arabic / English)
    ".office-preview img{max-width:100%;height:auto;}"
)


def show_office_preview(doc_id: str, d: dict):
    ext = ext_of(d["name"])

    # 1) Exact look: render pages through LibreOffice -> PDF -> PyMuPDF (when LibreOffice is installed)
    if ext != "html" and fitz is not None:
        pdf = office_to_pdf(doc_id, d["name"], d["data"])
        if pdf:
            try:
                imgs, total = page_images(f"{doc_id}-pdf", "preview.pdf", pdf)
                for i, img in enumerate(imgs, 1):
                    st.image(img, caption=f"Page {i} of {total}")
                if total > len(imgs):
                    st.caption(f"Preview limited to the first {len(imgs)} pages (the whole file is analyzed).")
                return
            except Exception:  # noqa: BLE001
                pass

    # 2) Fallback: simplified rendering from the file's content
    body = office_html(d["name"], d["data"])
    if body:
        st.html(f'<style>{_PREVIEW_CSS}</style><div class="office-preview">{body}</div>')  # st.html sanitizes the markup
        if ext != "html":
            st.caption("Simplified preview (layout may differ from the original). Install LibreOffice for page-accurate previews.")
        return
    st.info(
        "Preview isn't available for this file. Install `mammoth` (Word), `openpyxl` (Excel) or "
        "`python-pptx` (PowerPoint) to enable it. The file can still be analyzed."
    )


def show_preview(doc_id: str, d: dict):
    if ext_of(d["name"]) in OFFICE_TYPES:
        show_office_preview(doc_id, d)
        return
    try:
        if fitz is None:
            raise RuntimeError("PyMuPDF not installed")
        imgs, total = page_images(doc_id, d["name"], d["data"])
        for i, img in enumerate(imgs, 1):
            st.image(img, caption=f"Page {i} of {total}")
        if total > len(imgs):
            st.caption(f"Preview limited to the first {len(imgs)} pages (the whole file is analyzed).")
    except Exception:  # noqa: BLE001
        try:
            st.image(d["data"])
        except Exception:  # noqa: BLE001
            st.info("Preview isn't available for this file type, but it can still be analyzed.")


def show_df(df: pd.DataFrame):
    col_cfg = {}
    if "Confidence" in df.columns and pd.api.types.is_numeric_dtype(df["Confidence"]):
        col_cfg["Confidence"] = st.column_config.ProgressColumn( "Confidence", min_value=0.0, max_value=1.0, format="%.2f" )
    st.dataframe(df, hide_index=True, column_config=col_cfg)


def render_results(res: dict, tab_names: list, doc_id: str, name: str, model_label: str):
    chips = "".join(f'<span class="chip">{html.escape(s)}</span>' for s in res["stats"])
    st.markdown(f'<div class="chips">{chips}</div>', unsafe_allow_html=True)

    shown = [t for t in tab_names if res["sections"].get(t)]
    labels = [f"{t} ({sum(len(df) for _, df in res['sections'][t])})" for t in shown]
    tabs = st.tabs(labels + ["Text", "JSON"])

    for tab, t in zip(tabs, shown):
        with tab:
            for title, df in res["sections"][t]:
                if title:
                    st.markdown(f"**{title}**")
                show_df(df)
    if not shown:
        st.warning("The model returned no structured data for this document. Check the Text tab.")

    slug = model_label.lower().replace(" ", "_")
    with tabs[-2]:
        st.text_area(
            "Extracted text", res["text"], height=300, label_visibility="collapsed",
            key=f"text_{doc_id}_{model_label}",
        )
    with tabs[-1]:
        st.download_button(
            "Download JSON", res["json"], file_name=f"{os.path.splitext(name)[0]}_{slug}.json",
            mime="application/json", key=f"dl_{doc_id}_{model_label}",
        )
        st.json(res["json"], expanded=False)


def empty_state(title: str, body: str):
    st.markdown(f'<div class="empty"><b>{html.escape(title)}</b>{html.escape(body)}</div>', unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# State
# ----------------------------------------------------------------------------
st.session_state.setdefault("docs", {})  # doc_id -> {name, data, size, results, errors}
st.session_state.setdefault("uploader_key", 0)
docs: dict = st.session_state.docs

if msg := st.session_state.pop("toast", None):
    st.toast(msg)

endpoint, api_key = ENV_ENDPOINT, ENV_KEY
creds_ok = bool(endpoint and api_key)


def reset_page():
    st.session_state.docs = {}
    st.session_state.uploader_key += 1
    st.session_state.pop("view_doc", None)


# ----------------------------------------------------------------------------
# Header + step tracker (filled in once state is known)
# ----------------------------------------------------------------------------
st.markdown('<p class="app-title">Azure Document Intelligence</p>', unsafe_allow_html=True)
st.markdown(
    '<p class="app-sub">Extract fields, text and tables from invoices, receipts and other documents.</p>',
    unsafe_allow_html=True,
)
stepper_slot = st.empty()

left, right = st.columns([2, 7], gap="large")

# ----------------------------------------------------------------------------
# Left: control panel (model + upload + queue)
# ----------------------------------------------------------------------------
with left:
    with st.container(key="controls"):
        st.markdown('<p class="panel-h">Model</p>', unsafe_allow_html=True)
        model_label = st.selectbox(
            "Prebuilt model", list(MODELS), label_visibility="collapsed", on_change=reset_page, key="model"
        )
        cfg = MODELS[model_label]
        st.markdown(f'<p class="hint"><code>{cfg["model_id"]}</code> · {html.escape(cfg["help"])}</p>', unsafe_allow_html=True)

        st.markdown('<p class="panel-h" style="margin-top:1.25rem">Documents</p>', unsafe_allow_html=True)
        allowed_types = ALLOWED_TYPES if cfg["office"] else [t for t in ALLOWED_TYPES if t not in OFFICE_TYPES]
        uploaded = st.file_uploader(
            "Upload documents", type=allowed_types, accept_multiple_files=True,
            key=f"uploader_{st.session_state.uploader_key}", label_visibility="collapsed",
        )
        pills = "".join(f'<span class="pill">{t.upper()}</span>' for t in allowed_types)
        st.markdown(f'<div class="pills"><span class="pill limit">200 MB per file</span>{pills}</div>', unsafe_allow_html=True)
        if not cfg["office"]:
            st.markdown(
                '<p class="hint">This model accepts PDF and images only. '
                'Use Read or Layout for Word, Excel, PowerPoint and HTML files.</p>',
                unsafe_allow_html=True,
            )

        if uploaded:
            new_ids = []
            for f in uploaded:
                data = f.getvalue()
                doc_id = hashlib.sha1(data).hexdigest()[:12]
                if doc_id not in docs:
                    docs[doc_id] = {"name": f.name, "data": data, "size": len(data), "results": {}, "errors": {}}
                    new_ids.append(doc_id)
            if new_ids:
                st.session_state.view_doc = new_ids[-1]
            dupes = len(uploaded) - len(new_ids)
            st.session_state.toast = f"Added {len(new_ids)} document(s)" + (f" ({dupes} duplicate skipped)" if dupes else "")
            st.session_state.uploader_key += 1
            st.rerun()

        progress_slot = st.empty()

        if docs:
            st.markdown(f'<p class="panel-h" style="margin-top:1.25rem">Queue ({len(docs)})</p>', unsafe_allow_html=True)
            for doc_id, d in list(docs.items()):
                if d["errors"].get(model_label):
                    state, label = "err", "Failed"
                elif model_label in d["results"]:
                    state, label = "done", "Analyzed"
                else:
                    state, label = "", "Ready to analyze"
                with st.container(border=True):
                    c_thumb, c_info, c_del = st.columns([2, 6, 1.2], vertical_alignment="center")
                    with c_thumb:
                        thumb = thumbnail(doc_id, d["name"], d["data"])
                        if thumb:
                            st.image(thumb)
                        else:
                            st.markdown(f'<div class="ext">{ext_of(d["name"]).upper()}</div>', unsafe_allow_html=True)
                    with c_info:
                        st.markdown(
                            f'<div class="qname" title="{html.escape(d["name"])}">{html.escape(d["name"])}</div>'
                            f'<div class="qmeta">{fmt_size(d["size"])} · {label}</div>'
                            f'<div class="bar"><i class="{state}"></i></div>',
                            unsafe_allow_html=True,
                        )
                    with c_del:
                        if st.button("✕", key=f"del_{doc_id}", help="Remove this document"):
                            del docs[doc_id]
                            st.rerun()
            if st.button("Remove all"):
                docs.clear()
                st.rerun()

        if not creds_ok:
            st.error(
                "Azure credentials not found. Set DOCUMENTINTELLIGENCE_ENDPOINT and "
                "DOCUMENTINTELLIGENCE_API_KEY in your .env file, then restart the app."
            )
        with st.expander("Options"):
            force = st.checkbox("Re-analyze documents that already have results for this model")

# ----------------------------------------------------------------------------
# Right: split view (original document | extracted data)
# ----------------------------------------------------------------------------
with right:
    with st.container(key="workspace"):
        if not docs:
            empty_state(
                "Your results will appear here",
                "Add a document on the left to see the original beside its extracted fields.",
            )
        else:
            if st.session_state.get("view_doc") not in docs:
                st.session_state.view_doc = next(iter(docs))
            if len(docs) > 1:
                st.selectbox("Viewing", list(docs), format_func=lambda i: docs[i]["name"], key="view_doc")
            view_id = st.session_state.view_doc
            d = docs[view_id]

            pv, rs = st.columns([4, 5], gap="medium")
            with pv:
                st.markdown('<p class="panel-h">Original</p>', unsafe_allow_html=True)
                with st.container(height=680):
                    show_preview(view_id, d)
            with rs:
                st.markdown(f'<p class="panel-h">Extracted data · {html.escape(model_label)}</p>', unsafe_allow_html=True)
                if err := d["errors"].get(model_label):
                    st.error(err)
                elif res := d["results"].get(model_label):
                    render_results(res, cfg["tabs"], view_id, d["name"], model_label)
                else:
                    empty_state(
                        "Not analyzed yet",
                        f"Use the Analyze button to extract data with {model_label}.",
                    )

# ----------------------------------------------------------------------------
# Step tracker state
# ----------------------------------------------------------------------------
if not docs:
    current = 2
elif not any(model_label in d["results"] for d in docs.values()):
    current = 3
else:
    current = 4
stepper_slot.markdown(stepper_html(current), unsafe_allow_html=True)

# ----------------------------------------------------------------------------
# Floating call-to-action: only exists while there is something to analyze
# ----------------------------------------------------------------------------
pending = [i for i, d in docs.items() if force or model_label not in d["results"]]
if pending:
    with st.container(key="cta"):
        n = len(pending)
        run = st.button(f"Analyze {n} document{'s' if n != 1 else ''}", type="primary", disabled=not creds_ok)

    if run:
        client = get_client(endpoint, api_key)

        def work(doc_id: str):
            return parse(analyze(client, cfg, docs[doc_id]["data"], docs[doc_id]["name"]))

        with progress_slot.container():
            bar = st.progress(0.0, text=f"Analyzing {n} document(s)...")
            with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
                futures = {pool.submit(work, i): i for i in pending}
                for k, fut in enumerate(as_completed(futures), 1):
                    doc_id = futures[fut]
                    d = docs[doc_id]
                    try:
                        d["results"][model_label] = fut.result()
                        d["errors"].pop(model_label, None)
                    except HttpResponseError as e:
                        d["errors"][model_label] = e.message or str(e)
                    except Exception as e:  # noqa: BLE001
                        d["errors"][model_label] = str(e)
                    bar.progress(k / n, text=f"Finished {d['name']} ({k}/{n})")
        progress_slot.empty()
        st.rerun()
